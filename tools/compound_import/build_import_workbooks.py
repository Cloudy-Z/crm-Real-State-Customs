"""Build Frappe v16 Data Import workbooks for Compounds, Phases and Amenities.

Source: the user's Nawy export (one row per Compound or Phase, amenities as
check-mark columns). Output: one single-sheet .xlsx per DocType, in the order
they must be imported, plus a JSON summary used by the verification step.

Frappe facts this relies on (checked in the v16 importer source):
* Only the ACTIVE sheet of an .xlsx is read, so every file has one sheet.
* Parent headers may be the field Label; child headers are "Label (Table Label)".
* Rows whose parent columns are all blank are child rows of the row above,
  which is how several Amenities are attached to one Compound.
* Link values must already exist, so masters are imported before Compounds,
  and Compounds before Phases.

Usage: python3 build_import_workbooks.py <source.xlsx> <output_dir>
"""

import collections
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AMENITY_FIRST_COLUMN = 10  # zero-based index of "Outdoor pools"
HEADER_FILL = PatternFill("solid", fgColor="DDE7F5")
CHILD_FILL = PatternFill("solid", fgColor="EAF5E4")

# Readable, consistent amenity names. Every source column is kept.
AMENITY_NAME_FIXES = {
    "Outdoor pools": "Outdoor Pools",
    "Commercial strip": "Commercial Strip",
    "Jogging trail": "Jogging Trail",
    "bicycles lanes": "Bicycle Lanes",
    "Underground parking": "Underground Parking",
    "Medical center": "Medical Center",
    "Disability support": "Disability Support",
}

# Same display name, different project (La Vista Developments vs STM).
# Frappe names are case-insensitive, so the second one needs a distinct name.
DUPLICATE_NAME_FIXES = {
    "https://www.nawy.com/compound/2003-la-vista-city": "La Vista City (STM Developments)",
}


def clean(value):
    return " ".join(str(value or "").split())


def is_checked(value):
    return clean(value) in {"✓", "✔", "1", "yes", "Yes", "TRUE", "True", "x", "X"}


def read_source(path):
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean(cell) for cell in rows[0]]
    records = [dict(zip(headers, row)) for row in rows[1:] if any(v not in (None, "") for v in row)]
    amenity_columns = headers[AMENITY_FIRST_COLUMN:]
    return headers, records, amenity_columns


def build_model(records, amenity_columns):
    amenity_names = {col: AMENITY_NAME_FIXES.get(col, col) for col in amenity_columns}
    names_by_url = {}
    compounds, phases, problems = [], [], []

    for record in records:
        url = clean(record.get("compound_url"))
        name = DUPLICATE_NAME_FIXES.get(url, clean(record.get("compound_name")))
        names_by_url[url] = name

    seen = collections.Counter(name.lower() for name in names_by_url.values())
    for name, count in seen.items():
        if count > 1:
            problems.append(f"Duplicate name after cleanup: {name}")

    for record in records:
        url = clean(record.get("compound_url"))
        kind = clean(record.get("record_type")).lower()
        item = {
            "name": names_by_url[url],
            "record_type": "Phase" if kind == "phase" else "Compound",
            "developer": clean(record.get("developer")),
            "destination": clean(record.get("area")),
            "master_plan": clean(record.get("master_plan_urls")) or None,
            "source_url": url or None,
            "amenities": [amenity_names[col] for col in amenity_columns if is_checked(record.get(col))],
        }
        if kind not in {"compound", "phase"}:
            problems.append(f"Unknown record_type {record.get('record_type')!r} for {item['name']}")
        if kind == "phase":
            parent_url = clean(record.get("parent_compound_url"))
            parent = names_by_url.get(parent_url) or clean(record.get("PARENT COMPOUND"))
            if parent_url not in names_by_url:
                problems.append(f"Phase {item['name']} has no matching parent row")
            item["parent_compound"] = parent
            phases.append(item)
        else:
            compounds.append(item)

    compound_names = {c["name"] for c in compounds}
    for phase in phases:
        if phase["parent_compound"] not in compound_names:
            problems.append(f"Phase {phase['name']} parent {phase['parent_compound']} is not a Compound")
    return compounds, phases, list(amenity_names.values()), problems


def write_sheet(path, title, headers, rows, child_columns=0, widths=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
        is_child = index > len(headers) - child_columns
        cell.fill = CHILD_FILL if is_child else HEADER_FILL
        width = (widths or {}).get(header, max(14, min(60, len(header) + 4)))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    book.active = 0
    book.save(path)


def project_rows(items, include_parent):
    """One row per record, then one extra row per additional Amenity."""
    rows = []
    for item in items:
        parent_values = [item["name"], item["record_type"]]
        if include_parent:
            parent_values.append(item["parent_compound"])
        parent_values += [item["developer"], item["destination"], item["master_plan"], item["source_url"]]
        amenities = item["amenities"] or [None]
        rows.append(parent_values + [amenities[0]])
        blank = [None] * len(parent_values)
        for amenity in amenities[1:]:
            rows.append(blank + [amenity])
    return rows


def main(source, output_dir):
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    _headers, records, amenity_columns = read_source(source)
    compounds, phases, amenities, problems = build_model(records, amenity_columns)
    if problems:
        raise SystemExit("Data problems:\n" + "\n".join(problems))

    destinations = sorted({item["destination"] for item in compounds + phases})
    developers = sorted({item["developer"] for item in compounds + phases}, key=str.lower)

    write_sheet(
        output / "1_Destinations.xlsx", "Destinations",
        ["Destination Name", "Active"], [[name, 1] for name in destinations],
        widths={"Destination Name": 28},
    )
    write_sheet(
        output / "2_Developers.xlsx", "Developers",
        ["Developer Name"], [[name] for name in developers],
        widths={"Developer Name": 45},
    )
    write_sheet(
        output / "3_Amenities.xlsx", "Amenities",
        ["Amenity Name"], [[name] for name in amenities],
        widths={"Amenity Name": 32},
    )
    widths = {
        "Compound Name": 40, "Record Type": 13, "Parent Compound": 36, "Developer": 32,
        "Destination": 22, "Master Plan": 60, "Source URL": 55, "Amenity (Amenities)": 28,
    }
    base = ["Compound Name", "Record Type"]
    tail = ["Developer", "Destination", "Master Plan", "Source URL", "Amenity (Amenities)"]
    compound_rows = project_rows(compounds, include_parent=False)
    phase_rows = project_rows(phases, include_parent=True)
    write_sheet(output / "4_Compounds.xlsx", "Compounds", base + tail, compound_rows, 1, widths)
    write_sheet(output / "5_Phases.xlsx", "Phases", base + ["Parent Compound"] + tail, phase_rows, 1, widths)

    summary = {
        "destinations": len(destinations),
        "developers": len(developers),
        "amenities": len(amenities),
        "compounds": len(compounds),
        "phases": len(phases),
        "compound_rows": len(compound_rows),
        "phase_rows": len(phase_rows),
        "compound_amenity_links": sum(len(c["amenities"]) for c in compounds),
        "phase_amenity_links": sum(len(p["amenities"]) for p in phases),
        "renamed": DUPLICATE_NAME_FIXES,
        "amenity_name_fixes": AMENITY_NAME_FIXES,
        "destination_names": destinations,
    }
    (output / "import_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False))
    print(json.dumps({k: v for k, v in summary.items() if isinstance(v, int)}, indent=2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
